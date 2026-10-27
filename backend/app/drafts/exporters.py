"""
Renderers: block list -> HTML / DOCX / PDF.

All three consume the same `blocks` from engine.generate(), so the preview,
the Word file and the PDF never disagree.

Formatting follows common Indian court-filing conventions: A4, Times New
Roman 12pt, 1.5 line spacing, a wider left margin for binding, and justified
body text.

Library choices:
* DOCX via python-docx: the standard pure-Python way to write real .docx
  (editable by advocates, which matters more than the PDF).
* PDF via xhtml2pdf (reportlab underneath): pure Python, no system packages.
  WeasyPrint looks nicer but needs GTK/Pango, which is painful on Windows and
  bloats the Docker image.
"""

from __future__ import annotations

import html
import io
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

# ------------------------------------------------------------------- HTML
_CSS = """
@page { size: A4; margin: 2.5cm 2.5cm 2.5cm 3.5cm; }
body { font-family: 'Times New Roman', Times, serif; font-size: 12pt; line-height: 1.5; color: #000; }
p { margin: 0 0 8pt 0; text-align: justify; }
.court { text-align: center; font-weight: bold; font-size: 13pt; margin-bottom: 2pt; }
.center { text-align: center; }
.title { text-align: center; font-weight: bold; text-decoration: underline; margin: 10pt 0; }
.heading { font-weight: bold; text-decoration: underline; margin-top: 10pt; }
.left { text-align: left; }
.right { text-align: right; }
.para { text-indent: 1.25cm; }
.num { padding-left: 1cm; text-indent: -1cm; }
.sub { padding-left: 2cm; text-indent: -1cm; }
.space { height: 10pt; }
.rule { border-bottom: 1px solid #000; margin: 4pt 0 10pt 0; }
.pagebreak { page-break-after: always; }
"""


def _block_html(b: dict[str, Any]) -> str:
    kind = b["kind"]
    if kind == "space":
        return '<p class="space">&nbsp;</p>'
    if kind == "pagebreak":
        return '<div class="pagebreak"></div>'
    if kind == "rule":
        return '<div class="rule"></div>'
    text = html.escape(b["text"])  # user input is always escaped (no HTML injection)
    if b.get("label"):
        text = f"{html.escape(b['label'])}&nbsp;&nbsp;{text}"
    return f'<p class="{html.escape(kind)}">{text}</p>'


def to_html(draft: dict[str, Any], *, standalone: bool = True) -> str:
    body = "\n".join(_block_html(b) for b in draft["blocks"])
    if not standalone:
        return body
    return (f"<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(draft['title'])}</title>"
            f"<style>{_CSS}</style></head><body>{body}</body></html>")


# ------------------------------------------------------------------- PDF
def to_pdf(draft: dict[str, Any]) -> bytes:
    from xhtml2pdf import pisa  # imported lazily: it's slow to import and only needed here

    # The built-in PDF fonts (Times/Helvetica) have no glyph for the rupee sign
    # (U+20B9), so it would print as a black box. Embedding a TTF just for one
    # glyph adds size and licensing questions; "Rs." is the established
    # alternative in Indian pleadings anyway. DOCX keeps the real ₹.
    pdf_draft = {**draft, "blocks": [{**b, "text": b["text"].replace("₹", "Rs. ")} for b in draft["blocks"]]}
    buf = io.BytesIO()
    result = pisa.CreatePDF(to_html(pdf_draft), dest=buf, encoding="utf-8")
    if result.err:
        raise RuntimeError("PDF rendering failed")
    return buf.getvalue()


# ------------------------------------------------------------------- DOCX
def _bottom_border(paragraph) -> None:
    """python-docx has no API for paragraph borders, so add the OOXML <w:pBdr> by hand."""
    p_pr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in {"w:val": "single", "w:sz": "6", "w:space": "1", "w:color": "000000"}.items():
        bottom.set(qn(k), v)
    bdr.append(bottom)
    p_pr.append(bdr)


def to_docx(draft: dict[str, Any]) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)  # A4, since python-docx defaults to US Letter
    sec.left_margin, sec.right_margin = Cm(3.5), Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    # East-Asian font slot must also be set, or Word substitutes a font for some glyphs.
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    align = {
        "court": WD_ALIGN_PARAGRAPH.CENTER, "center": WD_ALIGN_PARAGRAPH.CENTER, "title": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT, "left": WD_ALIGN_PARAGRAPH.LEFT, "heading": WD_ALIGN_PARAGRAPH.LEFT,
    }
    for b in draft["blocks"]:
        kind = b["kind"]
        if kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            continue
        p = doc.add_paragraph()
        if kind == "space":
            continue
        if kind == "rule":
            _bottom_border(p)
            continue
        p.alignment = align.get(kind, WD_ALIGN_PARAGRAPH.JUSTIFY)
        fmt = p.paragraph_format
        if kind == "para":
            fmt.first_line_indent = Cm(1.25)
        elif kind in ("num", "sub"):
            # Hanging indent: label in the margin, wrapped lines align with the text.
            base = Cm(1.0) if kind == "num" else Cm(2.0)
            fmt.left_indent = base
            fmt.first_line_indent = -Cm(1.0)
            fmt.tab_stops.add_tab_stop(base)
        text = f"{b['label']}\t{b['text']}" if b.get("label") else b["text"]
        run = p.add_run(text)
        if kind in ("court", "title", "heading"):
            run.bold = True
        if kind in ("title", "heading"):
            run.underline = True
        if kind == "court":
            run.font.size = Pt(13)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
